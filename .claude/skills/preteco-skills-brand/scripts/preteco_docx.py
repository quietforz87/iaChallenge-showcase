"""Helpers para generar .docx con la marca Preteco a partir de assets/template.dotx.

Uso:
    import sys; sys.path.insert(0, "<ruta-a-la-skill>/scripts")
    from preteco_docx import new_document, add_table, add_ficha, add_callout, add_code, add_caption, add_toc

    doc = new_document()
    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph("Texto…")
    doc.save("salida.docx")

Todo el formato sale de los estilos nativos del template; las funciones solo aplican
estilos y, en tablas, refuerzan el relleno por celda para que LibreOffice/PDF lo respete.
"""
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATE = SKILL_DIR / "assets" / "template.dotx"

DOC_DARK = "334000"
LIMA = "CBFF00"
TABLE_LABEL = "EAF2C2"
TABLE_BAND = "F3F7DE"
WHITE = "FFFFFF"
USABLE_WIDTH_CM = 15.59

_DOTX_CT = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
_DOCX_CT = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def new_document(template: Path = TEMPLATE):
    """Abre el .dotx como documento nuevo (python-docx no abre .dotx directamente) y vacía el cuerpo."""
    buf = io.BytesIO()
    with zipfile.ZipFile(template) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(_DOTX_CT, _DOCX_CT)
            zout.writestr(item, data)
    buf.seek(0)
    doc = Document(buf)
    body = doc.element.body
    for p in list(body.iterchildren(qn("w:p"))):
        body.remove(p)
    return doc


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:shd"))
    if old is not None:
        tcpr.remove(old)
    tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def _set_widths(table, widths_cm):
    """Anchos por columna (grid + layout fijo, para que Word y LibreOffice los respeten),
    alto mínimo de fila 0,7 cm, filas que no se parten y centrado vertical (spec §5.1)."""
    table.autofit = False
    tw = [round(w / 2.54 * 1440) for w in widths_cm]
    tblpr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        old = tblpr.find(qn(tag))
        if old is not None:
            tblpr.remove(old)
    tblpr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{sum(tw)}" w:type="dxa"/>'))
    tblpr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), tw):
        gc.set(qn("w:w"), str(w))
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))  # una fila nunca se parte entre páginas
        trpr.append(parse_xml(f'<w:trHeight {nsdecls("w")} w:val="397" w:hRule="atLeast"/>'))
        seen = {}
        for cell, w in zip(row.cells, widths_cm):
            seen[id(cell._tc)] = (cell, seen.get(id(cell._tc), (cell, 0))[1] + w)
        for cell, w in seen.values():
            cell.width = Cm(w)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _header_run_format(cell):
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = None
            rpr = r._r.get_or_add_rPr()
            rpr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="{LIMA}"/>'))


def _mark_header_row(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def _keep_together(table):
    """Mantiene la tabla en una página si entra: todas las filas menos la última con keep-with-next."""
    for row in table.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True


def add_table(doc, header, rows, widths_cm=None, keep_together=True):
    """Tabla de contenido (variante B): header #334000/lima, 1ª columna bold, bandas blanco/#F3F7DE.
    keep_together=True evita que una tabla corta quede partida (usar False en tablas de más de una página)."""
    ncols = len(header)
    t = doc.add_table(rows=1 + len(rows), cols=ncols, style="Table Grid")
    widths_cm = widths_cm or [USABLE_WIDTH_CM / ncols] * ncols
    for i, text in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = str(text)
        _shade(c, DOC_DARK)
        _header_run_format(c)
    _mark_header_row(t.rows[0])
    for r_i, values in enumerate(rows, start=1):
        fill = WHITE if r_i % 2 == 1 else TABLE_BAND
        for c_i, text in enumerate(values):
            c = t.rows[r_i].cells[c_i]
            c.text = str(text)
            _shade(c, fill)
            if c_i == 0:
                for r in c.paragraphs[0].runs:
                    r.bold = True
    _set_widths(t, widths_cm)
    if keep_together:
        _keep_together(t)
    return t


def add_ficha(doc, title, pairs, label_ratio=0.36):
    """Ficha de datos (variante A): header combinado centrado en MAYÚSCULAS; etiqueta #EAF2C2 bold, valor blanco."""
    t = doc.add_table(rows=1 + len(pairs), cols=2, style="Table Grid")
    head = t.rows[0].cells[0].merge(t.rows[0].cells[1])
    head.text = title.upper()
    head.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _shade(head, DOC_DARK)
    _header_run_format(head)
    _mark_header_row(t.rows[0])
    for i, (label, value) in enumerate(pairs, start=1):
        lc, vc = t.rows[i].cells
        lc.text, vc.text = str(label), str(value)
        _shade(lc, TABLE_LABEL)
        _shade(vc, WHITE)
        for r in lc.paragraphs[0].runs:
            r.bold = True
    w = USABLE_WIDTH_CM
    _set_widths(t, [w * label_ratio, w * (1 - label_ratio)])
    _keep_together(t)
    return t


def add_callout(doc, text, label=None):
    """Callout = estilo nativo Intense Quote (fondo #F3F7DE, borde izquierdo 3 pt lima)."""
    p = doc.add_paragraph(style="Intense Quote")
    if label:
        r = p.add_run(f"{label}: ")
        r.bold = True
        r._r.get_or_add_rPr().append(parse_xml(f'<w:color {nsdecls("w")} w:val="{DOC_DARK}"/>'))
    p.add_run(text)
    return p


def add_code(doc, code: str):
    """Bloque de código = estilo nativo HTML Preformatted (DM Mono 9 pt, fondo #F3F7DE)."""
    lines = code.rstrip("\n").split("\n")
    last = None
    for line in lines:
        last = doc.add_paragraph(line, style="HTML Preformatted")
    if last is not None:
        last.paragraph_format.space_after = Cm(0.28)
    return last


def add_caption(doc, kind, text):
    """Caption nativo con numeración automática: kind = 'Figura' o 'Tabla'."""
    p = doc.add_paragraph(style="Caption")
    p.add_run(f"{kind} ")
    for xml in (
        f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>',
        f'<w:instrText {nsdecls("w")} xml:space="preserve"> SEQ {kind} \\* ARABIC </w:instrText>',
        f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>',
    ):
        p.add_run()._r.append(parse_xml(xml))
    p.add_run("1")
    p.add_run()._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
    p.add_run(f". {text}")
    if kind == "Figura":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_toc(doc, title="Índice"):
    """Índice (solo si > 6 páginas o > 5 Heading 1). Título con TOC Heading; campo TOC sobre Heading 1-2."""
    doc.add_paragraph(title, style="TOC Heading")
    p = doc.add_paragraph()
    for xml in (
        f'<w:fldChar {nsdecls("w")} w:fldCharType="begin" w:dirty="true"/>',
        f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText>',
        f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>',
    ):
        p.add_run()._r.append(parse_xml(xml))
    p.add_run("Actualizá el índice: clic derecho → Actualizar campo.")
    p.add_run()._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
    return p


def restart_numbering(doc, paragraph):
    """Reinicia una List Number en 1 (python-docx continúa la numeración entre listas)."""
    numbering = doc.part.numbering_part.element
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    new_id = max(ids) + 1
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{new_id}"><w:abstractNumId w:val="91"/>'
        f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
    ))
    ppr = paragraph._p.get_or_add_pPr()
    ppr.insert(0, parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{new_id}"/></w:numPr>'
    ))
    return new_id
